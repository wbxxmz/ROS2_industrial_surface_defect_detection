"""阶段 7 §7 验收：质检报告模块单元测试。

覆盖：字段一致性、空批次异常、无写权限异常、图片缺失降级、批量资源释放。
"""

import logging
import os
import warnings
from pathlib import Path

import pytest

from core.types import Detection
from core.verdict import Verdict
from judge import summarize
from report import ReportBuilder


def make_detection(class_id=0, severity=1, confidence=0.5):
    """构造一条符合阶段 0 契约的 Detection。"""
    return Detection(
        class_id=class_id,
        class_name={0: "scratch", 1: "dent", 2: "stain", 3: "notch"}[class_id],
        bbox=(10, 20, 30, 40),
        area_px=1200,
        area_ratio=0.03,
        confidence=confidence,
        mask=None,
        severity=severity,
    )


def make_verdict(is_pass=False, max_severity=1, detections=None):
    """构造与 detections 口径一致的 Verdict。"""
    detections = detections or []
    class_counts = {0: 0, 1: 0, 2: 0, 3: 0}
    for det in detections:
        class_counts[det.class_id] += 1
    return Verdict(
        is_pass=is_pass,
        reason="PASS" if is_pass else "检测到缺陷",
        total_count=len(detections),
        class_counts=class_counts,
        total_area_px=sum(det.area_px for det in detections),
        total_area_ratio=min(1.0, sum(det.area_ratio for det in detections)),
        triggered=[] if is_pass else ["count_exceeded"],
        max_severity=max_severity,
    )


def make_batch(image_path=None, count=2):
    """构造批量记录：每份样本 1 条 scratch 缺陷。"""
    batch = []
    for index in range(count):
        detections = [make_detection(confidence=0.5 + index * 0.1)]
        batch.append(
            {
                "source": f"sample_{index:03d}",
                "verdict": make_verdict(detections=detections),
                "detections": detections,
                "image_path": image_path,
            }
        )
    return batch


def base_cfg(tmp_path, fmt="xlsx"):
    return {
        "output_path": str(tmp_path / f"report.{fmt}"),
        "format": fmt,
        "title": "测试报告",
        "algorithm_name": "OpenCV",
        "algorithm_version": "0.1.0",
        "params": {"min_area": 10},
    }


def test_consistency_with_summarize(tmp_path):
    """报告字段与 summarize() / Verdict 逐字段一致，max_severity=None 展示为未分级。"""
    from openpyxl import load_workbook

    batch = make_batch(count=3)
    batch[0]["verdict"] = make_verdict(detections=batch[0]["detections"], max_severity=None)

    output = ReportBuilder().build(batch, base_cfg(tmp_path))
    assert output.is_file() and output.stat().st_size > 0

    summary = summarize([item["verdict"] for item in batch])
    workbook = load_workbook(output)
    sheet = workbook["汇总"]
    values = {
        str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    }

    assert values["样本总数"] == summary["sample_count"]
    assert values["合格数"] == summary["pass_count"]
    assert values["不合格数"] == summary["fail_count"]
    assert values["良率"] == f"{summary['yield_rate'] * 100:.2f}%"
    assert values["缺陷总数"] == summary["total_count"]
    assert values["缺陷总面积占比"] == f"{summary['total_area_ratio'] * 100:.2f}%"
    assert values["批次最高严重度"] == "中度"  # 其余样本 max_severity=1
    # 类别计数与 summarize() 一致（按"缺陷类别"表头定位，避免脚注行影响行号）
    class_header_row = next(
        row for row in range(1, sheet.max_row + 1) if sheet.cell(row=row, column=1).value == "缺陷类别"
    )
    assert sheet.cell(row=class_header_row + 1, column=2).value == summary["class_counts"][0]
    assert "未达置信度阈值" in str(sheet.cell(row=sheet.max_row, column=1).value)
    # 明细行数与缺陷总数一致
    assert workbook["缺陷明细"].max_row - 1 == summary["total_count"]


def test_consistency_none_severity_is_not_zero(tmp_path):
    """max_severity=None 必须展示为"未分级"，不得当作 0（轻度）。"""
    from openpyxl import load_workbook

    batch = make_batch(count=1)
    batch[0]["verdict"] = make_verdict(detections=batch[0]["detections"], max_severity=None)

    output = ReportBuilder().build(batch, base_cfg(tmp_path))
    sheet = load_workbook(output)["汇总"]
    values = {
        str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value
    }
    assert values["批次最高严重度"] == "未分级"


def test_counted_flag_matches_judge_prefilter(tmp_path):
    """明细表「计入判定」与 judge 预筛同口径：低于置信度阈值标「否」。"""
    from openpyxl import load_workbook

    cfg = base_cfg(tmp_path)
    cfg["judge"] = {"min_confidence": 0.25, "min_area_px": 10}
    valid = make_detection(confidence=0.50)
    low_conf = make_detection(confidence=0.10)
    batch = [
        {
            "source": "sample_low",
            "verdict": make_verdict(detections=[valid]),
            "detections": [valid, low_conf],
            "image_path": None,
        }
    ]

    output = ReportBuilder().build(batch, cfg)
    rows = list(load_workbook(output)["缺陷明细"].iter_rows(min_row=2, values_only=True))
    assert [row[-1] for row in rows] == ["是", "否"]


def test_empty_batch_raises_and_creates_nothing(tmp_path):
    """空批次抛 ValueError，且不落盘任何报告文件（§4.4）。"""
    cfg = base_cfg(tmp_path)
    with pytest.raises(ValueError):
        ReportBuilder().build([], cfg)
    assert not Path(cfg["output_path"]).exists()

    with pytest.raises(ValueError):
        ReportBuilder().build(None, cfg)
    assert not Path(cfg["output_path"]).exists()


def test_invalid_item_raises(tmp_path):
    """元素缺必填键 / verdict 类型不符 → ValueError。"""
    cfg = base_cfg(tmp_path)
    with pytest.raises(ValueError):
        ReportBuilder().build([{"source": "x"}], cfg)
    with pytest.raises(ValueError):
        ReportBuilder().build(
            [{"source": "x", "verdict": "not-a-verdict", "detections": []}], cfg
        )


@pytest.mark.skipif(os.geteuid() == 0, reason="root 不受目录写权限限制")
def test_permission_error(tmp_path):
    """目标目录只读 → OSError（底层异常原样上抛）。"""
    readonly = tmp_path / "readonly"
    readonly.mkdir()
    readonly.chmod(0o555)
    try:
        cfg = {"output_path": str(readonly / "report.pdf"), "format": "pdf"}
        with pytest.raises(OSError):
            ReportBuilder().build(make_batch(), cfg)
    finally:
        readonly.chmod(0o755)


def test_missing_image_degrades(tmp_path, caplog):
    """image_path=None 或文件不存在 → 仍生成报告，且记 warning（§4.4）。"""
    cfg = base_cfg(tmp_path, fmt="pdf")
    batch = make_batch(image_path=None, count=1)
    batch.append(
        {
            "source": "missing",
            "verdict": make_verdict(detections=[make_detection()]),
            "detections": [make_detection()],
            "image_path": str(tmp_path / "not_exists.jpg"),
        }
    )
    with caplog.at_level(logging.WARNING):
        output = ReportBuilder().build(batch, cfg)

    assert output.is_file() and output.stat().st_size > 0
    assert any("缺失" in record.message or "missing" in record.message for record in caplog.records)


def test_batch_release(tmp_path):
    """连续 build 多份，文件句柄不随批次数单调增长，且无 ResourceWarning（§4.5）。"""
    builder = ReportBuilder()
    cfg = base_cfg(tmp_path, fmt="pdf")
    fd_counts = []
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        for _ in range(5):
            builder.build(make_batch(count=1), cfg)
            fd_counts.append(len(os.listdir("/proc/self/fd")))

    assert max(fd_counts) - min(fd_counts) <= 3
    assert not [item for item in caught if issubclass(item.category, ResourceWarning)]


def test_pdf_and_xlsx_output(tmp_path):
    """PDF / Excel 均可正确落盘（返回 Path 且大小 > 0）。"""
    batch = make_batch(count=2)
    for fmt in ("pdf", "xlsx"):
        output = ReportBuilder().build(batch, base_cfg(tmp_path, fmt=fmt))
        assert isinstance(output, Path)
        assert output.is_file() and output.stat().st_size > 0
