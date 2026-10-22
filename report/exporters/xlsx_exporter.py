"""阶段 7 §4.2：Excel 导出（openpyxl，无头落盘）。

三张工作表：汇总 / 缺陷明细 / 样本明细；类别分布图 PNG 嵌入汇总页。
``wb.save()`` 内部自行关闭写入流（§4.5），无需手动关闭。
"""

import logging
from pathlib import Path

logger = logging.getLogger("defect.report")


def export_xlsx(data: dict, output: Path) -> Path:
    """把组装好的数据写成 xlsx，返回路径（§5）。"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from report.charts import class_distribution_chart

    summary = data["summary"]
    workbook = Workbook()

    # ---------- 汇总 ----------
    sheet = workbook.active
    sheet.title = "汇总"
    sheet["A1"] = data["title"]
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.merge_cells("A1:D1")

    rows = [
        ("检测批次时间", data["generated_at"]),
        ("样本总数", summary["sample_count"]),
        ("合格数", summary["pass_count"]),
        ("不合格数", summary["fail_count"]),
        ("良率", f"{summary['yield_rate'] * 100:.2f}%"),
        ("缺陷总数", summary["total_count"]),
        ("缺陷总面积(px)", summary["total_area_px"]),
        ("缺陷总面积占比", f"{summary['total_area_ratio'] * 100:.2f}%"),
        ("批次最高严重度", data["overall_max_severity"]),
        ("算法", f"{data['algorithm_name']} {data['algorithm_version']}"),
        ("样本来源", data["sample_source"] or "未提供"),
        ("参数", str(data["params"]) if data["params"] is not None else "未提供"),
    ]
    for offset, (key, value) in enumerate(rows, start=3):
        sheet.cell(row=offset, column=1, value=key).font = Font(bold=True)
        sheet.cell(row=offset, column=2, value=value)

    class_row = 3 + len(rows) + 1
    sheet.cell(row=class_row, column=1, value="缺陷类别").font = Font(bold=True)
    sheet.cell(row=class_row, column=2, value="数量").font = Font(bold=True)
    for index, (class_id, class_name) in enumerate(zip(sorted(summary["class_counts"]), data["class_names"])):
        sheet.cell(row=class_row + 1 + index, column=1, value=class_name)
        sheet.cell(row=class_row + 1 + index, column=2, value=summary["class_counts"][class_id])

    note_row = class_row + 2 + len(data["class_names"])
    note = sheet.cell(
        row=note_row,
        column=1,
        value="* 明细包含所有原始检出，其中未达置信度阈值的框不计入最终判定。",
    )
    note.font = Font(size=9, italic=True, color="808080")

    chart_path = output.parent / f"{output.stem}_class_chart.png"
    class_distribution_chart(summary, chart_path, data["class_names"])
    image = XLImage(str(chart_path))
    image.width, image.height = 480, 224
    sheet.add_image(image, "D3")
    temp_files = [chart_path]  # 统计图在 wb.save() 时才真正读取，保存后再清理

    # ---------- 缺陷明细 ----------
    detail_sheet = workbook.create_sheet("缺陷明细")
    detail_header = [
        "样本", "类别", "类别ID", "位置(x,y,w,h)", "面积(px)", "面积占比", "置信度", "严重度", "计入判定",
    ]
    detail_sheet.append(detail_header)
    for row in data["details"]:
        detail_sheet.append(
            [
                row["sample"],
                row["class_name"],
                row["class_id"],
                row["bbox"],
                row["area_px"],
                row["area_ratio"],
                row["confidence"],
                row["severity"],
                row["counted"],
            ]
        )
    for cell in detail_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="DCE3EC")
        cell.font = Font(bold=True)
    for offset, row in enumerate(data["details"], start=2):
        if not row["counted_raw"]:  # 未计入判定：整行浅灰，与有效缺陷在视觉上区分
            for cell in detail_sheet[offset]:
                cell.font = Font(color="808080")

    # ---------- 样本明细 ----------
    sample_sheet = workbook.create_sheet("样本明细")
    sample_sheet.append(["样本", "判定", "判定原因", "缺陷数", "最高严重度", "标注图路径"])
    for row in data["samples"]:
        sample_sheet.append(
            [row["source"], row["result"], row["reason"], row["total_count"],
             row["max_severity"], row["image_path"] or "N/A"]
        )
    for cell in sample_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="DCE3EC")
        cell.font = Font(bold=True)

    for target, widths in (
        (sheet, (22, 40, 4, 4)),
        (detail_sheet, (26, 12, 8, 20, 10, 12, 10, 10, 12)),
        (sample_sheet, (26, 10, 40, 10, 14, 60)),
    ):
        for index, width in enumerate(widths, start=1):
            target.column_dimensions[get_column_letter(index)].width = width
        for row_cells in target.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(vertical="center")

    try:
        workbook.save(output)  # 内部关闭写入流（§4.5）
    finally:
        for temp_file in temp_files:
            temp_file.unlink(missing_ok=True)  # 临时统计图不残留
    return output
